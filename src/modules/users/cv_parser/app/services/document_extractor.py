"""Document text extraction for PDF and DOCX CVs (fully in-memory)."""

import io
from dataclasses import dataclass, field
from typing import List

import docx
from pypdf import PdfReader


class ExtractionError(Exception):
    pass


@dataclass
class Page:
    page: int
    text: str


@dataclass
class ExtractedDocument:
    filename: str
    file_type: str
    pages: List[Page] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)


def extract_pdf(content: bytes, filename: str) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [
            Page(page=i + 1, text=(page.extract_text() or ""))
            for i, page in enumerate(reader.pages)
        ]
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("PDF could not be read") from exc

    if not any(p.text.strip() for p in pages):
        raise ExtractionError("No text could be extracted from the PDF (it may be a scanned image).")
    return ExtractedDocument(filename=filename, file_type="pdf", pages=pages)


def extract_docx(content: bytes, filename: str) -> ExtractedDocument:
    try:
        document = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in document.paragraphs]
        # Include table text (CVs sometimes use tables for layout).
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        text = "\n".join(paragraphs)
    except Exception as exc:
        raise ExtractionError("DOCX could not be read") from exc

    if not text.strip():
        raise ExtractionError("No text could be extracted from the DOCX.")
    # DOCX has no page structure; treat the whole body as one page.
    return ExtractedDocument(
        filename=filename, file_type="docx", pages=[Page(page=1, text=text)]
    )


def extract_document(content: bytes, file_type: str, filename: str) -> ExtractedDocument:
    if file_type == "pdf":
        return extract_pdf(content, filename)
    if file_type == "docx":
        return extract_docx(content, filename)
    raise ExtractionError(f"Unsupported file type: {file_type}")
